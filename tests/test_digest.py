"""The workbook's one unforgivable failure would be eating her work.

Everything else here is presentation. The read-back tests are the product.
"""

import datetime
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from report.digest import (  # noqa: E402
    HER_COLUMNS,
    SHEET_BIDS,
    SHEET_HEALTH,
    SHEET_PIPELINE,
    SHEET_TODAY,
    build,
    days_left,
    read_back_her_columns,
)

TODAY = datetime.date(2026, 9, 19)


def _cand(**kw):
    base = dict(
        pncp_key="46316600000164-1-000447/2025",
        municipio="Itaquaquecetuba", uf="SP",
        objeto="AQUISICAO DE KIT MATERNIDADE",
        valor_estimado=2380250.00,
        encerramento="2026-10-10T09:00:00",
        quantidade=5000,
    )
    base.update(kw)
    return base


class _Health:
    rows_total, rows_served, total_expected, pages_fetched = 4595, 4595, 4595, 93
    rows_duplicate, missing, route_mixed, failures = 0, 0, False, []
    complete = True
    suspected_outage = False


# -- the thing that must never break ---------------------------------------

def test_her_typed_columns_survive_a_rebuild(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)

    # she fills in her side
    wb = load_workbook(path)
    ws = wb[SHEET_BIDS]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("Eu licitei?") + 1).value = "SIM"
    ws.cell(row=2, column=headers.index("Meu lance (R$)") + 1).value = 357.04
    ws.cell(row=2, column=headers.index("Resultado") + 1).value = "PERDI"
    ws.cell(row=2, column=headers.index("Preço vencedor (R$)") + 1).value = 349.00
    ws.cell(row=2, column=headers.index("Minhas anotações") + 1).value = "frete alto"
    wb.save(path)

    # the machine rebuilds the next morning with fresh data
    build(path, [_cand(valor_estimado=999.0)], health=_Health(), today=TODAY)

    ws = load_workbook(path)[SHEET_BIDS]
    headers = [c.value for c in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(headers)}
    assert row["Eu licitei?"] == "SIM"
    assert row["Meu lance (R$)"] == 357.04
    assert row["Resultado"] == "PERDI"
    assert row["Preço vencedor (R$)"] == 349.00
    assert row["Minhas anotações"] == "frete alto"


def test_a_rebuild_keeps_a_dated_backup(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)
    build(path, [_cand()], health=_Health(), today=TODAY)
    backups = os.listdir(tmp_path / "arquivo")
    assert any(b.startswith("2026-09-19") for b in backups)


def test_read_back_survives_a_missing_or_corrupt_file(tmp_path):
    assert read_back_her_columns(str(tmp_path / "nope.xlsx")) == {}
    junk = tmp_path / "junk.xlsx"
    junk.write_text("not a workbook")
    assert read_back_her_columns(str(junk)) == {}


def test_the_machine_never_writes_her_columns_itself(tmp_path):
    """A rebuild with no prior file must leave her columns EMPTY, not guessed."""
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)
    ws = load_workbook(path)[SHEET_BIDS]
    headers = [c.value for c in ws[1]]
    for column in HER_COLUMNS:
        # openpyxl stores an empty string as None; either way the machine must
        # not have put a value there. What matters is that it guessed nothing.
        assert not ws.cell(row=2, column=headers.index(column) + 1).value


# -- structure -------------------------------------------------------------

def test_all_four_sheets_exist_and_are_in_portuguese(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)
    wb = load_workbook(path)
    assert wb.sheetnames == [SHEET_TODAY, SHEET_PIPELINE, SHEET_BIDS, SHEET_HEALTH]
    headers = [c.value for c in wb[SHEET_TODAY][1]]
    assert "Município" in headers and "DECISÃO" in headers
    assert "Dias restantes" in headers


def test_rows_are_ordered_by_urgency(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [
        _cand(pncp_key="a/2026", municipio="Longe", encerramento="2026-12-01T09:00:00"),
        _cand(pncp_key="b/2026", municipio="Urgente", encerramento="2026-09-22T09:00:00"),
    ], health=_Health(), today=TODAY)
    ws = load_workbook(path)[SHEET_TODAY]
    assert ws.cell(row=2, column=1).value == "Urgente"


def test_days_left_handles_junk():
    assert days_left("2026-09-25T09:00:00", TODAY) == 6
    assert days_left(None, TODAY) is None
    assert days_left("sem data", TODAY) is None


# -- honesty ---------------------------------------------------------------

def test_an_unpriceable_tender_says_VERIFICAR_not_NAO_LICITAR(tmp_path):
    """A tender dropped for a reason nobody stated is a tender silently lost."""
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)
    ws = load_workbook(path)[SHEET_TODAY]
    headers = [c.value for c in ws[1]]
    verdict = ws.cell(row=2, column=headers.index("DECISÃO") + 1).value
    why = ws.cell(row=2, column=headers.index("Por quê") + 1).value
    assert verdict == "VERIFICAR"
    assert "itens" in why


def test_a_failed_rule_produces_NAO_LICITAR_with_its_reason(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand(rule_results={
        "6": {"passed": False, "is_flag": False,
              "reason": "comprador ainda deve 99% das faturas do ano anterior"}})],
        health=_Health(), today=TODAY)
    ws = load_workbook(path)[SHEET_TODAY]
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("DECISÃO") + 1).value == "NAO LICITAR"
    assert "99%" in ws.cell(row=2, column=headers.index("Por quê") + 1).value


def test_an_unscreened_buyer_is_shown_as_unverifiable(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)
    ws = load_workbook(path)[SHEET_TODAY]
    headers = [c.value for c in ws[1]]
    assert (ws.cell(row=2, column=headers.index("O comprador paga?") + 1).value
            == "NÃO VERIFICÁVEL")


def test_a_short_harvest_is_stated_loudly(tmp_path):
    class Short(_Health):
        complete = False
        missing = 300
        rows_total = 4295

    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=Short(), today=TODAY)
    ws = load_workbook(path)[SHEET_HEALTH]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
    assert "FALHOU" in text
    assert "INCOMPLETA" in text


def test_a_missing_health_report_is_itself_a_warning(tmp_path):
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=None, today=TODAY)
    ws = load_workbook(path)[SHEET_HEALTH]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
    assert "ATENÇÃO" in text
    assert "completa" in text


def test_the_recall_ceiling_is_always_stated(tmp_path):
    """She must never read the day's count as the whole market."""
    path = str(tmp_path / "ACOLHE.xlsx")
    build(path, [_cand()], health=_Health(), today=TODAY)
    ws = load_workbook(path)[SHEET_HEALTH]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
    assert "1 em cada 4" in text
