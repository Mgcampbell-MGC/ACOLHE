"""The daily entrypoint, end to end, with a stubbed client and no network.

What this pins: the steps run in order, an UNAVAILABLE descent reaches the
workbook as a precise VERIFICAR rather than crashing or vanishing, and a
harvest failure is never turned into a quiet day.
"""

import datetime
import json
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from harvest.pncp_client import Unavailable  # noqa: E402
from harvest.run_daily import run  # noqa: E402
from report.digest import SHEET_HEALTH, SHEET_TODAY  # noqa: E402

DAY = "20260915"
TODAY = datetime.date(2026, 9, 19)
NOW = datetime.datetime(2026, 9, 19, 12, 0)


def _tender(key, mun, uf, objeto, closes="2026-10-10T09:00:00", valor=100000.0):
    return {
        "numeroControlePNCP": key,
        "objetoCompra": objeto,
        "valorTotalEstimado": valor,
        "dataPublicacaoPncp": "2026-09-15T10:00:00",
        "dataEncerramentoProposta": closes,
        "orgaoEntidade": {"cnpj": key[:14], "razaoSocial": f"MUNICIPIO DE {mun}",
                          "esferaId": "M"},
        "unidadeOrgao": {"municipioNome": mun, "ufSigla": uf, "codigoIbge": "0"},
    }


class _Report:
    def __init__(self, n):
        self.rows_total = self.rows_served = self.total_expected = n
        self.pages_fetched, self.rows_duplicate, self.missing = 1, 0, 0
        self.route_mixed, self.failures = False, []

    @property
    def complete(self):
        return True

    @property
    def suspected_outage(self):
        return False


class StubClient:
    """harvest_day writes rows into the sink; itens behaves as configured."""

    def __init__(self, rows, itens_mode="unavailable", harvest_raises=False):
        self.rows = rows
        self.itens_mode = itens_mode
        self.harvest_raises = harvest_raises
        self.itens_calls = 0

    def harvest_day(self, day, modalidades, sink):
        if self.harvest_raises:
            raise Unavailable("https://pncp.gov.br/api/consulta/v1/x", [])
        for row in self.rows:
            sink(row)
        return _Report(len(self.rows))

    def itens(self, cnpj, ano, sequencial, pagina=1, page_size=50):
        self.itens_calls += 1
        if self.itens_mode == "unavailable":
            raise Unavailable("https://pncp.gov.br/api/pncp/v1/x", [])
        if self.itens_mode == "kit":
            return {"data": [
                {"numeroItem": 1, "descricao": "Body manga longa 100% algodao RN",
                 "quantidade": 200, "unidadeMedida": "UN"},
                {"numeroItem": 2, "descricao": "Mamadeira 50 ml bico silicone",
                 "quantidade": 200, "unidadeMedida": "UN"},
            ], "totalPaginas": 1}
        return {"data": [], "totalPaginas": 1}


ROWS = [
    _tender("46316600000164-1-000447/2025", "Itaquaquecetuba", "SP",
            "AQUISICAO DE KIT MATERNIDADE"),
    _tender("11111111000111-1-000001/2026", "Maracas", "BA",
            "Fornecimento de Kit Enxoval de bebe"),
    # A NEAR MISS must first look like a kit ('KIT ENXOVAL') and then be
    # excluded ('ESCOLAR'). 'KIT ESCOLAR' alone never matches a kit pattern,
    # so it is merely unrelated -- the same fixture mistake test_daily.py
    # already corrected once today.
    _tender("22222222000122-1-000002/2026", "Escola", "MT",
            "AQUISICAO DE KIT ENXOVAL ESCOLAR PARA A REDE MUNICIPAL"),
    _tender("33333333000133-1-000003/2026", "Velho", "PR",
            "KIT NATALIDADE", closes="2026-09-18T23:59:59"),  # already closed
    _tender("44444444000144-1-000004/2026", "Asfalto", "GO",
            "PAVIMENTACAO ASFALTICA"),
]


def _run(tmp_path, client, **kw):
    out = str(tmp_path / "ACOLHE.xlsx")
    res = run(DAY, out, client=client, state_dir=str(tmp_path / "state"),
              cache_dir=str(tmp_path / "cache"), today=TODAY, now=NOW, **kw)
    return res, out


def test_the_pipeline_runs_in_order_and_writes_the_workbook(tmp_path):
    res, out = _run(tmp_path, StubClient(ROWS))
    assert res["rows"] == 5
    assert {c["municipio"] for c in res["candidates"]} == {"Itaquaquecetuba", "Maracas"}
    assert len(res["near_misses"]) == 1 and "ESCOLAR" in res["near_misses"][0]["excluded_by"]
    assert len(res["stale"]) == 1 and res["stale"][0]["municipio"] == "Velho"
    assert os.path.exists(out)
    # the harvest sink is on disk, one JSON row per line
    sink = tmp_path / "state" / f"harvest-{DAY}.jsonl"
    assert len([l for l in open(sink) if l.strip()]) == 5


def test_an_unavailable_descent_reaches_the_workbook_as_a_precise_verificar(tmp_path):
    """Family B down is a fact about PNCP, not a crash and not an empty kit."""
    client = StubClient(ROWS, itens_mode="unavailable")
    res, out = _run(tmp_path, client)
    assert client.itens_calls == 2                      # one per candidate
    assert all(c["items_status"] == "UNAVAILABLE" for c in res["candidates"])
    ws = load_workbook(out)[SHEET_TODAY]
    headers = [c.value for c in ws[1]]
    reasons = [ws.cell(row=r, column=headers.index("Por quê") + 1).value
               for r in range(2, ws.max_row + 1)]
    assert all("indisponivel" in (why or "") for why in reasons)
    assert all("Nao e um kit vazio" in (why or "") for why in reasons)


def test_a_readable_kit_is_classified_and_carried(tmp_path):
    client = StubClient(ROWS, itens_mode="kit")
    res, _ = _run(tmp_path, client)
    cand = res["candidates"][0]
    assert cand["items_status"] == "OK"
    skus = {i["numero"]: i["sku"] for i in cand["items"]}
    assert skus[1] == "BODY"
    assert skus[2] is None                                # mamadeira refused


def test_no_descend_flag_skips_family_b_entirely(tmp_path):
    client = StubClient(ROWS, itens_mode="unavailable")
    res, _ = _run(tmp_path, client, do_descend=False)
    assert client.itens_calls == 0
    assert res["descent"] == []
    assert "items_status" not in res["candidates"][0]


def test_a_harvest_failure_raises_and_writes_no_workbook(tmp_path):
    """A listing outage must never become a quiet day with an empty sheet."""
    out = str(tmp_path / "ACOLHE.xlsx")
    with pytest.raises(Unavailable):
        run(DAY, out, client=StubClient(ROWS, harvest_raises=True),
            state_dir=str(tmp_path / "state"), cache_dir=str(tmp_path / "cache"),
            today=TODAY, now=NOW)
    assert not os.path.exists(out)


def test_health_sheet_reports_the_dropped_rows(tmp_path):
    _, out = _run(tmp_path, StubClient(ROWS))
    ws = load_workbook(out)[SHEET_HEALTH]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row)
    assert "fora do prazo" in text          # the closed PR tender
    assert "outro programa" in text         # the school-uniform tender
